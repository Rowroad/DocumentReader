"""Format converters for exporting documents to various formats."""

import os
import json
from pathlib import Path
from typing import List, Optional
from datetime import datetime
import xml.etree.ElementTree as ET
from xml.dom import minidom

from docx import Document as DocxDocument
from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from ebooklib import epub
import markdown
from pydub import AudioSegment
from mutagen.mp4 import MP4, MP4Cover

from ..models import Document, OutputFormat, Settings
from .gemini_client import GeminiClient
from .tts_manager import TTSManager


class FormatConverter:
    """Converts documents to various output formats."""

    def __init__(self, settings: Settings, gemini_client: Optional[GeminiClient] = None):
        """Initialize the format converter.

        Args:
            settings: Application settings
            gemini_client: Gemini API client for TTS (optional)
        """
        self.settings = settings
        self.gemini = gemini_client
        self.tts_manager = TTSManager(settings, gemini_client)

    def convert(self, document: Document, output_format: OutputFormat,
                output_path: Path) -> Path:
        """Convert a document to the specified format.

        Args:
            document: Document to convert
            output_format: Target format
            output_path: Output file path

        Returns:
            Path to the created file
        """
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        if output_format == OutputFormat.TXT:
            return self._to_txt(document, output_path)
        elif output_format == OutputFormat.HTML:
            return self._to_html(document, output_path)
        elif output_format == OutputFormat.MARKDOWN:
            return self._to_markdown(document, output_path)
        elif output_format == OutputFormat.DOCX:
            return self._to_docx(document, output_path)
        elif output_format == OutputFormat.EPUB:
            return self._to_epub(document, output_path)
        elif output_format == OutputFormat.PDF:
            return self._to_pdf(document, output_path)
        elif output_format == OutputFormat.DAISY_TEXT:
            return self._to_daisy_text(document, output_path)
        elif output_format == OutputFormat.DAISY_AUDIO:
            return self._to_daisy_audio(document, output_path)
        elif output_format == OutputFormat.M4B:
            return self._to_m4b(document, output_path)
        else:
            raise ValueError(f"Unsupported output format: {output_format}")

    def _to_txt(self, document: Document, output_path: Path) -> Path:
        """Convert to plain text."""
        with open(output_path, 'w', encoding='utf-8') as f:
            # Title
            f.write(f"{document.title}\n")
            f.write("=" * len(document.title) + "\n\n")

            # Metadata
            if document.metadata:
                if 'author' in document.metadata:
                    f.write(f"Author: {document.metadata['author']}\n")
                if 'description' in document.metadata:
                    f.write(f"Description: {document.metadata['description']}\n")
                f.write("\n")

            # Content
            f.write(document.content)

        return output_path

    def _to_html(self, document: Document, output_path: Path) -> Path:
        """Convert to accessible HTML."""
        html_content = f"""<!DOCTYPE html>
<html lang="{document.language}">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{document.title}</title>
    <style>
        body {{
            font-family: Arial, sans-serif;
            line-height: 1.6;
            max-width: 800px;
            margin: 0 auto;
            padding: 20px;
        }}
        h1, h2, h3, h4, h5, h6 {{
            margin-top: 1.5em;
            margin-bottom: 0.5em;
        }}
        img {{
            max-width: 100%;
            height: auto;
        }}
        table {{
            border-collapse: collapse;
            width: 100%;
            margin: 1em 0;
        }}
        th, td {{
            border: 1px solid #ddd;
            padding: 8px;
            text-align: left;
        }}
        th {{
            background-color: #f2f2f2;
        }}
    </style>
</head>
<body>
    <main role="main">
        <h1>{document.title}</h1>
"""

        # Add metadata
        if document.metadata:
            html_content += "        <section aria-label='Document metadata'>\n"
            if 'author' in document.metadata:
                html_content += f"            <p><strong>Author:</strong> {document.metadata['author']}</p>\n"
            if 'description' in document.metadata:
                html_content += f"            <p><strong>Description:</strong> {document.metadata['description']}</p>\n"
            html_content += "        </section>\n"

        # Add table of contents if structure exists
        if document.structure:
            html_content += self._generate_toc_html(document.structure)

        # Add content with structure
        html_content += self._structure_to_html(document.structure, document.content)

        # Add images
        if document.images:
            html_content += "        <section aria-label='Images'>\n"
            for img in document.images:
                if img.path:
                    html_content += f"            <figure>\n"
                    html_content += f"                <img src='{img.path}' alt='{img.alt_text}' />\n"
                    if img.ocr_text:
                        html_content += f"                <figcaption>Text: {img.ocr_text}</figcaption>\n"
                    html_content += f"            </figure>\n"
            html_content += "        </section>\n"

        html_content += """    </main>
</body>
</html>"""

        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(html_content)

        return output_path

    def _to_markdown(self, document: Document, output_path: Path) -> Path:
        """Convert to Markdown."""
        md_content = f"# {document.title}\n\n"

        # Metadata
        if document.metadata:
            if 'author' in document.metadata:
                md_content += f"**Author:** {document.metadata['author']}\n\n"
            if 'description' in document.metadata:
                md_content += f"**Description:** {document.metadata['description']}\n\n"

        # Structure
        if document.structure:
            md_content += "## Table of Contents\n\n"
            md_content += self._generate_toc_markdown(document.structure)
            md_content += "\n\n"
            md_content += self._structure_to_markdown(document.structure)
        else:
            md_content += document.content

        # Images
        if document.images:
            md_content += "\n\n## Images\n\n"
            for img in document.images:
                if img.path:
                    md_content += f"![{img.alt_text}]({img.path})\n\n"
                if img.ocr_text:
                    md_content += f"*Text from image: {img.ocr_text}*\n\n"

        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(md_content)

        return output_path

    def _to_docx(self, document: Document, output_path: Path) -> Path:
        """Convert to DOCX."""
        doc = DocxDocument()

        # Title
        title = doc.add_heading(document.title, level=0)
        title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        # Metadata
        if document.metadata:
            if 'author' in document.metadata:
                p = doc.add_paragraph()
                p.add_run(f"Author: {document.metadata['author']}").italic = True
            if 'description' in document.metadata:
                p = doc.add_paragraph()
                p.add_run(f"Description: {document.metadata['description']}").italic = True

        doc.add_paragraph()  # Spacing

        # Add structure or content
        if document.structure:
            self._add_structure_to_docx(doc, document.structure)
        else:
            doc.add_paragraph(document.content)

        doc.save(output_path)
        return output_path

    def _to_epub(self, document: Document, output_path: Path) -> Path:
        """Convert to EPUB."""
        book = epub.EpubBook()

        # Metadata
        book.set_identifier(f'id{document.created_at.timestamp()}')
        book.set_title(document.title)
        book.set_language(document.language)

        if document.metadata.get('author'):
            book.add_author(document.metadata['author'])

        # Create chapters from structure
        chapters = []
        spine = ['nav']

        if document.structure:
            for idx, section in enumerate(document.structure):
                chapter = epub.EpubHtml(
                    title=section.title,
                    file_name=f'chap_{idx:03d}.xhtml',
                    lang=document.language
                )
                chapter.content = f'<h1>{section.title}</h1><p>{section.content}</p>'
                book.add_item(chapter)
                chapters.append(chapter)
                spine.append(chapter)
        else:
            # Single chapter
            chapter = epub.EpubHtml(
                title=document.title,
                file_name='content.xhtml',
                lang=document.language
            )
            chapter.content = f'<h1>{document.title}</h1><p>{document.content}</p>'
            book.add_item(chapter)
            chapters.append(chapter)
            spine.append(chapter)

        # Add navigation
        book.toc = chapters
        book.add_item(epub.EpubNcx())
        book.add_item(epub.EpubNav())

        # Spine
        book.spine = spine

        # Write
        epub.write_epub(output_path, book)
        return output_path

    def _to_pdf(self, document: Document, output_path: Path) -> Path:
        """Convert to accessible PDF (via HTML intermediate)."""
        # Create HTML first
        html_path = output_path.with_suffix('.html')
        self._to_html(document, html_path)

        # Convert HTML to PDF using external tool (requires wkhtmltopdf or similar)
        try:
            import pdfkit
            pdfkit.from_file(str(html_path), str(output_path))
            html_path.unlink()  # Remove temporary HTML
            return output_path
        except ImportError:
            raise NotImplementedError(
                "PDF conversion requires pdfkit and wkhtmltopdf. "
                "HTML file saved instead at: " + str(html_path)
            )

    def _to_daisy_text(self, document: Document, output_path: Path) -> Path:
        """Convert to DAISY text-only format."""
        # Create DAISY directory structure
        daisy_dir = output_path.parent / output_path.stem
        daisy_dir.mkdir(parents=True, exist_ok=True)

        # Create NCC (Navigation Control Center) file
        ncc_path = daisy_dir / "ncc.html"
        ncc_content = self._generate_daisy_ncc(document, audio=False)

        with open(ncc_path, 'w', encoding='utf-8') as f:
            f.write(ncc_content)

        # Create content SMIL files (text-only)
        if document.structure:
            for idx, section in enumerate(document.structure):
                smil_path = daisy_dir / f"content_{idx:03d}.smil"
                html_path = daisy_dir / f"content_{idx:03d}.html"

                # SMIL file
                smil_content = self._generate_text_smil(section, idx)
                with open(smil_path, 'w', encoding='utf-8') as f:
                    f.write(smil_content)

                # HTML file
                html_content = f"""<!DOCTYPE html>
<html><head><title>{section.title}</title></head>
<body><h1>{section.title}</h1><p>{section.content}</p></body></html>"""
                with open(html_path, 'w', encoding='utf-8') as f:
                    f.write(html_content)

        return ncc_path

    def _to_daisy_audio(self, document: Document, output_path: Path) -> Path:
        """Convert to DAISY audio + text format."""
        # Create DAISY directory structure
        daisy_dir = output_path.parent / output_path.stem
        daisy_dir.mkdir(parents=True, exist_ok=True)

        # Generate audio for each section
        audio_files = []

        if document.structure:
            for idx, section in enumerate(document.structure):
                audio_path = daisy_dir / f"audio_{idx:03d}.mp3"

                try:
                    # Generate TTS audio
                    self.tts_manager.generate_audio(
                        section.content,
                        audio_path,
                        section.title
                    )
                    audio_files.append((section, audio_path))
                except NotImplementedError as e:
                    # TTS not available, skip audio
                    print(f"Warning: {e}")
                    return self._to_daisy_text(document, output_path)

        # Create NCC file
        ncc_path = daisy_dir / "ncc.html"
        ncc_content = self._generate_daisy_ncc(document, audio=True)

        with open(ncc_path, 'w', encoding='utf-8') as f:
            f.write(ncc_content)

        # Create SMIL files with audio sync
        for idx, (section, audio_path) in enumerate(audio_files):
            smil_path = daisy_dir / f"content_{idx:03d}.smil"
            html_path = daisy_dir / f"content_{idx:03d}.html"

            # SMIL file with audio
            smil_content = self._generate_audio_smil(section, idx, audio_path.name)
            with open(smil_path, 'w', encoding='utf-8') as f:
                f.write(smil_content)

            # HTML file
            html_content = f"""<!DOCTYPE html>
<html><head><title>{section.title}</title></head>
<body><h1>{section.title}</h1><p>{section.content}</p></body></html>"""
            with open(html_path, 'w', encoding='utf-8') as f:
                f.write(html_content)

        return ncc_path

    def _to_m4b(self, document: Document, output_path: Path) -> Path:
        """Convert to M4B audiobook format."""
        temp_dir = Path(self.settings.temp_directory) / "m4b_temp"
        temp_dir.mkdir(parents=True, exist_ok=True)

        audio_segments = []
        chapter_markers = []
        current_time = 0

        # Generate audio for each section
        if document.structure:
            for idx, section in enumerate(document.structure):
                audio_path = temp_dir / f"chapter_{idx:03d}.mp3"

                try:
                    # Generate TTS audio
                    self.tts_manager.generate_audio(
                        section.content,
                        audio_path,
                        section.title
                    )

                    # Load audio segment
                    segment = AudioSegment.from_mp3(audio_path)
                    audio_segments.append(segment)

                    # Add chapter marker
                    chapter_markers.append({
                        'title': section.title,
                        'start_time': current_time
                    })

                    current_time += len(segment)

                except NotImplementedError as e:
                    raise Exception(f"TTS required for M4B generation: {e}")

        # Concatenate all audio
        combined = sum(audio_segments)

        # Export as M4A
        m4a_path = temp_dir / "temp.m4a"
        combined.export(m4a_path, format='ipod', codec='aac')

        # Add metadata and chapters
        audio = MP4(m4a_path)
        audio['\xa9nam'] = document.title  # Title

        if document.metadata.get('author'):
            audio['\xa9ART'] = document.metadata['author']  # Artist

        # Add chapters (requires mutagen support)
        # Note: Chapter support in M4B requires specific MP4 atoms

        audio.save()

        # Rename to .m4b
        m4a_path.rename(output_path)

        # Cleanup
        for file in temp_dir.glob('*.mp3'):
            file.unlink()

        return output_path

    # Helper methods

    def _generate_toc_html(self, structure: List) -> str:
        """Generate HTML table of contents."""
        html = "        <nav aria-label='Table of contents'>\n"
        html += "            <h2>Table of Contents</h2>\n"
        html += "            <ul>\n"

        for section in structure:
            html += f"                <li><a href='#{section.id}'>{section.title}</a></li>\n"
            if section.children:
                html += self._generate_toc_html_recursive(section.children, level=1)

        html += "            </ul>\n"
        html += "        </nav>\n"
        return html

    def _generate_toc_html_recursive(self, structure: List, level: int) -> str:
        """Recursively generate nested TOC."""
        html = "                <ul>\n"
        indent = "    " * (level + 4)

        for section in structure:
            html += f"{indent}<li><a href='#{section.id}'>{section.title}</a></li>\n"
            if section.children:
                html += self._generate_toc_html_recursive(section.children, level + 1)

        html += "                </ul>\n"
        return html

    def _generate_toc_markdown(self, structure: List, level: int = 0) -> str:
        """Generate Markdown table of contents."""
        md = ""
        indent = "  " * level

        for section in structure:
            md += f"{indent}- [{section.title}](#{section.id})\n"
            if section.children:
                md += self._generate_toc_markdown(section.children, level + 1)

        return md

    def _structure_to_html(self, structure: List, content: str) -> str:
        """Convert document structure to HTML."""
        if not structure:
            return f"        <div>{content}</div>\n"

        html = ""
        for section in structure:
            html += f"        <section id='{section.id}'>\n"
            html += f"            <h{section.level}>{section.title}</h{section.level}>\n"
            html += f"            <div>{section.content}</div>\n"

            if section.children:
                html += self._structure_to_html(section.children, "")

            html += "        </section>\n"

        return html

    def _structure_to_markdown(self, structure: List, level: int = 1) -> str:
        """Convert document structure to Markdown."""
        md = ""

        for section in structure:
            heading = "#" * (level + 1)
            md += f"{heading} {section.title}\n\n"
            md += f"{section.content}\n\n"

            if section.children:
                md += self._structure_to_markdown(section.children, level + 1)

        return md

    def _add_structure_to_docx(self, doc: DocxDocument, structure: List, level: int = 1):
        """Add document structure to DOCX."""
        for section in structure:
            doc.add_heading(section.title, level=level)
            doc.add_paragraph(section.content)

            if section.children:
                self._add_structure_to_docx(doc, section.children, level + 1)

    def _generate_daisy_ncc(self, document: Document, audio: bool = False) -> str:
        """Generate DAISY NCC file."""
        return f"""<!DOCTYPE html PUBLIC "-//W3C//DTD XHTML 1.0 Transitional//EN"
"http://www.w3.org/TR/xhtml1/DTD/xhtml1-transitional.dtd">
<html xmlns="http://www.w3.org/1999/xhtml">
<head>
<title>{document.title}</title>
<meta name="ncc:generator" content="DocumentReader"/>
<meta name="ncc:format" content="Daisy 2.02"/>
</head>
<body>
<h1 class="title">{document.title}</h1>
</body>
</html>"""

    def _generate_text_smil(self, section, idx: int) -> str:
        """Generate text-only SMIL file."""
        return f"""<?xml version="1.0"?>
<smil>
<head><meta name="ncc:generator" content="DocumentReader"/></head>
<body>
<seq>
<text src="content_{idx:03d}.html#text_{idx:03d}" id="text_{idx:03d}"/>
</seq>
</body>
</smil>"""

    def _generate_audio_smil(self, section, idx: int, audio_file: str) -> str:
        """Generate audio SMIL file."""
        return f"""<?xml version="1.0"?>
<smil>
<head><meta name="ncc:generator" content="DocumentReader"/></head>
<body>
<seq>
<par>
<text src="content_{idx:03d}.html#text_{idx:03d}" id="text_{idx:03d}"/>
<audio src="{audio_file}" clip-begin="0s"/>
</par>
</seq>
</body>
</smil>"""
